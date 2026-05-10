from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

from app.models.school import AcademicYear
from app.models.structure import Class, Subject
from app.reports.marks_excel import parse_subject_header, read_rows, write_template
from app.repositories import class_repo, school_repo


def _seed(conn: sqlite3.Connection) -> int:
    yid = school_repo.create_academic_year(
        conn,
        AcademicYear(
            id=None,
            label="2025-26",
            start_date="2025-04-01",
            end_date="2026-03-31",
            is_active=True,
        ),
    )
    cid = class_repo.create_class(conn, Class(id=None, name="5", section="A", academic_year_id=yid))
    class_repo.create_subject(conn, Subject(id=None, name="Math", class_id=cid, max_marks=100))
    class_repo.create_subject(conn, Subject(id=None, name="English", class_id=cid, max_marks=80))
    for i in range(2):
        conn.execute(
            "INSERT INTO students (admission_no, first_name, admission_date, class_id, status, roll_no) "
            "VALUES (?, ?, ?, ?, 'active', ?)",
            (f"ADM/{i + 1:03d}", f"Stu{i + 1:02d}", "2025-04-01", cid, i + 1),
        )
    return cid


def test_template_includes_subjects_and_students(tmp_path: Path, conn: sqlite3.Connection) -> None:
    cid = _seed(conn)
    out = write_template(tmp_path / "marks.xlsx", conn, class_id=cid, exam_name="Mid-term")
    assert out.is_file()
    rows = read_rows(out)
    # Two students, no marks filled in yet.
    assert len(rows) == 2
    assert rows[0]["admission_no"] == "ADM/001"
    # Subject headers carry the max-marks suffix.
    headers = list(rows[0]["marks"].keys())
    assert any("Math" in h and "max=100" in h for h in headers)
    assert any("English" in h and "max=80" in h for h in headers)


def test_parse_subject_header() -> None:
    assert parse_subject_header("Math (max=100)") == ("Math", 100)
    assert parse_subject_header("English (max=80)") == ("English", 80)
    # Missing max -> None
    assert parse_subject_header("Hindi") == ("Hindi", None)


def test_read_rows_handles_blank_cells(tmp_path: Path, conn: sqlite3.Connection) -> None:
    from openpyxl import load_workbook

    cid = _seed(conn)
    target = tmp_path / "marks.xlsx"
    write_template(target, conn, class_id=cid, exam_name="Test")
    wb = load_workbook(target)
    ws = wb.active
    # Fill row 2 with marks for one subject only.
    ws.cell(row=2, column=4, value=85)
    wb.save(target)

    rows = read_rows(target)
    first = rows[0]
    values = list(first["marks"].values())
    assert any(v == 85.0 for v in values)
    # The other subject column is still blank.
    assert any(v is None for v in values)
