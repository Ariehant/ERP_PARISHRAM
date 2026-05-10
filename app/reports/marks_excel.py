"""Marks Excel template + export + reader.

The template is per-class+exam. Columns are::

    Admission No | Roll | Name | <Subject1> (max=N) | <Subject2> (max=N) | ...

The import flow looks students up by ``admission_no`` (admission_no is
unique school-wide, roll_no can repeat between classes).
"""

from __future__ import annotations

import logging
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.repositories import class_repo, student_repo

log = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _subject_header(name: str, max_marks: int) -> str:
    return f"{name} (max={max_marks})"


def write_template(
    target: str | Path,
    conn: sqlite3.Connection,
    *,
    class_id: int,
    exam_name: str,
) -> Path:
    """Write a marks-entry template for the given class.

    Pre-fills admission no., roll, name; leaves subject columns blank.
    """
    target = Path(target)
    subjects = class_repo.list_subjects_for_class(conn, class_id)
    students = student_repo.list_all_for_export(conn, class_id=class_id, status="active")

    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"

    headers: list[str] = ["Admission No", "Roll", "Name"] + [
        _subject_header(s.name, s.max_marks) for s in subjects
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
    # Reasonable widths.
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 22
    for i in range(len(subjects)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 14

    for r, s in enumerate(students, start=2):
        ws.cell(row=r, column=1, value=s.admission_no)
        ws.cell(row=r, column=2, value=s.roll_no or "")
        ws.cell(row=r, column=3, value=" ".join(filter(None, [s.first_name, s.last_name])))

    # Notes sheet.
    notes = wb.create_sheet("Notes")
    notes["A1"] = f"Marks template for: {exam_name}"
    notes["A1"].font = Font(bold=True, size=14)
    note_lines = [
        "1. Do NOT change the 'Admission No' column -- it identifies the student.",
        "2. Each subject column shows its max marks in the header.",
        "3. Leave a cell blank to skip that student/subject. Cells > max_marks "
        "or < 0 will be rejected during import.",
        "4. Grades are computed automatically from the configured grade scale.",
    ]
    for i, line in enumerate(note_lines, start=3):
        notes.cell(row=i, column=1, value=line)
    notes.column_dimensions["A"].width = 90

    ws.freeze_panes = "D2"
    wb.save(target)
    log.info("Wrote marks template for class=%s to %s", class_id, target)
    return target


def read_rows(source: str | Path) -> list[dict[str, Any]]:
    """Read a marks file. Returns one dict per data row::

        {
            "admission_no": str,
            "marks": {subject_header: float | None},
        }

    Empty rows are skipped. Numeric cells parse as float; blank/whitespace
    cells become ``None``.
    """
    wb = load_workbook(source, data_only=True)
    ws = wb.worksheets[0]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    if len(headers) < 4:
        return []
    if not headers[0].lower().startswith("admission"):
        raise ValueError("First column must be 'Admission No'. Was the template edited?")
    subject_headers = headers[3:]

    out: list[dict[str, Any]] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in raw):
            continue
        admission = raw[0]
        if admission is None:
            continue
        admission_str = str(admission).strip()
        if not admission_str:
            continue
        marks: dict[str, float | None] = {}
        for h, value in zip(subject_headers, raw[3:], strict=False):
            if value is None or (isinstance(value, str) and not value.strip()):
                marks[h] = None
                continue
            try:
                marks[h] = float(value)
            except (TypeError, ValueError):
                marks[h] = None
        out.append({"admission_no": admission_str, "marks": marks})
    return out


def parse_subject_header(header: str) -> tuple[str, int | None]:
    """Pull the subject name and max marks out of a "Math (max=100)" header."""
    if "(max=" not in header:
        return header.strip(), None
    name, _, rest = header.partition("(max=")
    rest = rest.rstrip(") ")
    try:
        return name.strip(), int(rest)
    except ValueError:
        return name.strip(), None
